"""Import verified historical rankings without changing the source workbooks.

Run from the repository root with the bundled Python that includes openpyxl:
    python scripts/import-historical-results.py --check
    python scripts/import-historical-results.py --write
"""

from __future__ import annotations

import argparse
from io import BytesIO
import json
from pathlib import Path
import re
import unicodedata
import warnings
from zipfile import ZipFile

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SOURCES = ROOT.parent / "排名文档"
CONFIG = {
    "2022": ("结社设计大赛-获奖名单.xlsx", "结社设计大赛", "秘闻", "设定"),
    "2023": ("神话历史设计大赛-参赛统计.xlsx", "最终结果", "逸闻", "真相"),
    "2024": ("神话地理设计大赛-参赛统计.xlsx", "最终结果", "游记", "方志"),
    "2025": ("神话风俗设计大赛-参赛统计.xlsx", "最终结果", "采风", "录异"),
}
EXCLUDED_2025_ROWS = {12, 24, 36, 37, 40, 47, 49, 60, 61}
EXPECTED_COUNTS = {"2022": 52, "2023": 54, "2024": 70, "2025": 58}
AUTHOR_REVEAL_RANK_LIMITS = {"2022": 20, "2023": 35, "2024": 33, "2025": 32}
# The original submission marks this work anonymous; the publication display was confirmed anonymous.
ANONYMOUS_PUBLICATION_TITLES = {"2025": {"星之彩档案"}}


def open_source(path: Path):
    if "风俗" not in path.name:
        return load_workbook(path, read_only=True, data_only=True)

    # The source archive has empty <fill/> elements that openpyxl cannot parse.
    # Replace those elements only in a memory copy of the ZIP, keeping the XLSX intact.
    buffer = BytesIO()
    with ZipFile(path) as source, ZipFile(buffer, "w") as repaired:
        for member in source.infolist():
            data = source.read(member)
            if member.filename == "xl/styles.xml":
                data = data.replace(b"<fill/>", b'<fill><patternFill patternType="none"/></fill>')
            repaired.writestr(member, data)
    buffer.seek(0)
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", message="Workbook contains no default style")
        return load_workbook(buffer, read_only=True, data_only=True)


def frontmatter_value(raw: str, key: str) -> str:
    match = re.search(rf"(?m)^{re.escape(key)}:\s*(.*)$", raw)
    if not match:
        return ""
    value = match.group(1).strip()
    if value.startswith('"') and value.endswith('"'):
        return json.loads(value)
    return value.strip("'")


def normalized(text: str) -> str:
    text = unicodedata.normalize("NFKC", text).casefold()
    text = re.sub(r"[《》〈〉‘’'\"“”「」『』\s·•—–－_.,，。:：!！?？()（）\[\]【】]", "", text)
    return text


def works_by_title(year: str) -> dict[str, list[dict]]:
    works: dict[str, list[dict]] = {}
    for path in sorted((ROOT / "content" / year / "works").glob("*.md")):
        raw = path.read_text(encoding="utf-8")
        title = frontmatter_value(raw, "title")
        author = frontmatter_value(raw, "author")
        if not title:
            raise ValueError(f"Missing title: {path}")
        works.setdefault(normalized(title), []).append(
            {"slug": path.stem, "title": title, "author": author}
        )
    return works


def find_work(title: str, author: str, works: dict[str, list[dict]]) -> dict | None:
    candidates = works.get(normalized(title), [])
    if len(candidates) == 1:
        return candidates[0]
    if candidates:
        author_matches = [item for item in candidates if normalized(item["author"]) == normalized(author)]
        if len(author_matches) == 1:
            return author_matches[0]
        raise ValueError(f"Ambiguous title/author: {title!r}, {author!r}: {candidates}")
    return None


def numeric(value, year: str, row_number: int, field: str) -> float:
    if not isinstance(value, (int, float)) or isinstance(value, bool):
        raise ValueError(f"{year} row {row_number}: {field} is not numeric: {value!r}")
    return float(value)


def import_year(year: str) -> dict:
    file_name, sheet_name, first_label, second_label = CONFIG[year]
    path = SOURCES / file_name
    workbook = open_source(path)
    sheet = workbook[sheet_name]
    works = works_by_title(year)
    rankings = []
    ranks: dict[int, float] = {}
    seen_works = set()
    linked_slugs = set()
    for row_number, row in enumerate(sheet.values, start=1):
        if row_number < 3 or (year == "2025" and row_number in EXCLUDED_2025_ROWS):
            continue
        if all(value is None for value in row[:6]):
            continue
        rank = int(numeric(row[0], year, row_number, "rank"))
        source_no_raw = numeric(row[1], year, row_number, "sourceNo")
        if not source_no_raw.is_integer():
            raise ValueError(f"{year} row {row_number}: non-integer sourceNo")
        source_no = int(source_no_raw)
        title, author = row[2], row[3]
        if isinstance(author, (int, float)) and not isinstance(author, bool):
            author = str(int(author)) if float(author).is_integer() else str(author)
        if not isinstance(title, str) or not title.strip() or not isinstance(author, str) or not author.strip():
            raise ValueError(f"{year} row {row_number}: missing title or author")
        work_key = (normalized(title), normalized(author))
        if work_key in seen_works:
            raise ValueError(f"{year} row {row_number}: duplicate work {title!r} / {author!r}")
        seen_works.add(work_key)
        overall = numeric(row[4], year, row_number, "overall")
        if rank in ranks and abs(ranks[rank] - overall) > 1e-8:
            raise ValueError(f"{year} rank {rank} repeats with a different score at row {row_number}")
        ranks[rank] = overall
        first = numeric(row[11], year, row_number, first_label)
        second = numeric(row[17], year, row_number, second_label)
        if abs(overall - (first + second) / 2) > 1e-8:
            raise ValueError(f"{year} row {row_number}: overall differs from dimension mean")
        award = row[5]
        if award is not None and not isinstance(award, str):
            raise ValueError(f"{year} row {row_number}: unexpected award value {award!r}")
        award = award.strip() if award and award.strip() else None
        matched_work = find_work(title, author, works)
        slug = matched_work["slug"] if matched_work else None
        if slug is not None:
            if slug in linked_slugs:
                raise ValueError(f"{year} row {row_number}: duplicate link to {slug}")
            linked_slugs.add(slug)
        entry = {
            "rank": rank,
            "sourceNo": source_no,
            "workSlug": slug,
            "title": title.strip(),
            "author": (
                "匿名"
                if title.strip() in ANONYMOUS_PUBLICATION_TITLES.get(year, set())
                or (rank > AUTHOR_REVEAL_RANK_LIMITS[year] and not award)
                else author.strip()
            ),
            "overall": round(overall, 6),
            "award": award,
            "dimensions": [
                {"label": first_label, "average": round(first, 6)},
                {"label": second_label, "average": round(second, 6)},
            ],
        }
        if abs(entry["overall"] - sum(d["average"] for d in entry["dimensions"]) / 2) > 1e-6:
            raise ValueError(f"{year} row {row_number}: rounded score differs from rounded dimension mean")
        rankings.append(entry)
    workbook.close()
    if len(rankings) != EXPECTED_COUNTS[year]:
        raise ValueError(f"{year}: expected {EXPECTED_COUNTS[year]} rankings, got {len(rankings)}")
    rankings.sort(key=lambda entry: entry["rank"])
    return {
        "schemaVersion": 1,
        "edition": year,
        "source": {"file": file_name, "sheet": sheet_name},
        "authorRevealRankLimit": AUTHOR_REVEAL_RANK_LIMITS[year],
        "rankings": rankings,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    action = parser.add_mutually_exclusive_group(required=True)
    action.add_argument("--check", action="store_true", help="Validate and print a summary only")
    action.add_argument("--write", action="store_true", help="Write the four results.json files")
    args = parser.parse_args()
    results = {year: import_year(year) for year in CONFIG}
    for year, result in results.items():
        unlinked = [entry for entry in result["rankings"] if entry["workSlug"] is None]
        ranks = [entry["rank"] for entry in result["rankings"]]
        tied_ranks = sorted({rank for rank in ranks if ranks.count(rank) > 1})
        print(f"{year}: {len(result['rankings'])} rankings, {len(unlinked)} unlinked, tied ranks {tied_ranks}")
        for entry in unlinked:
            print(f"  rank {entry['rank']}: {entry['title']} / {entry['author']}")
        for entry in result["rankings"]:
            if entry["workSlug"] is None or entry["author"] == "匿名":
                continue
            work_path = ROOT / "content" / year / "works" / (entry["workSlug"] + ".md")
            archive_author = frontmatter_value(work_path.read_text(encoding="utf-8"), "author")
            if normalized(entry["author"]) != normalized(archive_author):
                print(f"  author differs, rank {entry['rank']}: source {entry['author']} / archive {archive_author}")
    if args.write:
        for year, result in results.items():
            path = ROOT / "content" / year / "results.json"
            path.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    else:
        for year, result in results.items():
            path = ROOT / "content" / year / "results.json"
            if path.exists() and json.loads(path.read_text(encoding="utf-8")) != result:
                raise ValueError(f"{path} differs from the verified source import")


if __name__ == "__main__":
    main()
